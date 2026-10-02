"""Turning fetched bytes into text: HTML pages and PDF, DOCX, XLSX, TXT, and CSV files."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath
from urllib.parse import urljoin, urlsplit, urlunsplit

from bs4 import BeautifulSoup
from charset_normalizer import from_bytes

MAX_UNCOMPRESSED_FACTOR = 20  # a zip may expand to at most 20x the download limit
MAX_COMPRESSION_RATIO = 100
MAX_SHEET_CELLS = 2_000_000

CONTENT_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "text/plain": ".txt",
    "text/csv": ".csv",
}


class ExtractionError(Exception):
    """A file could not be turned into text."""


@dataclass
class HtmlPage:
    title: str
    description: str
    text: str
    links: list[str] = field(default_factory=list)
    noindex: bool = False
    nofollow: bool = False


def normalize_url(url: str) -> str:
    """Lower-case scheme and host, drop default ports and fragments, keep path and query."""
    parts = urlsplit(url.strip())
    scheme = parts.scheme.lower()
    host = (parts.hostname or "").lower()
    port = parts.port
    if port and not ((scheme == "http" and port == 80) or (scheme == "https" and port == 443)):
        host = f"{host}:{port}"
    return urlunsplit((scheme, host, parts.path or "/", parts.query, ""))


def file_type(url: str, content_type: str = "") -> str:
    """The file extension for a URL, falling back to its content type."""
    suffix = PurePosixPath(urlsplit(url).path).suffix.lower()
    return suffix or CONTENT_TYPES.get(content_type, "")


def robots_directives(value: str) -> tuple[bool, bool]:
    """Parse a meta robots or X-Robots-Tag value into (noindex, nofollow)."""
    tokens = {token.strip().lower() for token in re.split(r"[,\s]+", value) if token.strip()}
    none = "none" in tokens
    return none or "noindex" in tokens, none or "nofollow" in tokens


def parse_html(content: bytes, url: str, encoding: str | None = None) -> HtmlPage:
    soup = BeautifulSoup(content, "html.parser", from_encoding=encoding)
    noindex = nofollow = False
    for meta in soup.find_all(
        "meta", attrs={"name": re.compile(r"^(robots|respectscraper)$", re.I)}
    ):
        page_noindex, page_nofollow = robots_directives(str(meta.get("content", "")))
        noindex = noindex or page_noindex
        nofollow = nofollow or page_nofollow

    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    description_tag = soup.find("meta", attrs={"name": re.compile(r"^description$", re.I)})
    description = str(description_tag.get("content", "")).strip() if description_tag else ""

    links: list[str] = []
    if not nofollow:
        seen: set[str] = set()
        for anchor in soup.find_all("a", href=True):
            rel = {value.lower() for value in (anchor.get("rel") or [])}
            if "nofollow" in rel:
                continue
            absolute = urljoin(url, str(anchor["href"]))
            if urlsplit(absolute).scheme not in ("http", "https"):
                continue
            normalized = normalize_url(absolute)
            if normalized not in seen:
                seen.add(normalized)
                links.append(normalized)

    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    body = soup.body or soup
    text = re.sub(r"\s+", " ", body.get_text(" ", strip=True)).strip()
    return HtmlPage(title, description, text, links, noindex, nofollow)


def decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        best = from_bytes(data).best()
        if best is None:
            raise ExtractionError("could not detect the text encoding") from None
        return str(best)


def extract_file(data: bytes, kind: str, *, max_bytes: int) -> str:
    """Extract text from a downloaded file. Raises ExtractionError on failure."""
    if kind in (".txt", ".csv"):
        return decode_text(data)
    if kind == ".pdf":
        return _pdf(data)
    if kind in (".docx", ".xlsx"):
        _check_zip(data, max_bytes)
        return _docx(data) if kind == ".docx" else _xlsx(data)
    raise ExtractionError(f"unsupported file type: {kind or 'unknown'}")


def _check_zip(data: bytes, max_bytes: int) -> None:
    """Refuse archives that would expand far beyond their size (zip bombs)."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ExtractionError("not a valid Office document") from exc
    total = 0
    for info in archive.infolist():
        total += info.file_size
        if info.compress_size and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO:
            raise ExtractionError(f"refused: {info.filename} has a suspicious compression ratio")
    if total > max_bytes * MAX_UNCOMPRESSED_FACTOR:
        raise ExtractionError("refused: the document expands beyond the size limit")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise ExtractionError(f"could not read PDF: {exc}") from exc
    return "\n\n".join(text.strip() for text in pages if text.strip())


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types
        raise ExtractionError(f"could not read DOCX: {exc}") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise ExtractionError(f"could not read XLSX: {exc}") from exc
    parts: list[str] = []
    cells = 0
    try:
        for sheet in workbook.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_SHEET_CELLS:
                    raise ExtractionError("refused: the workbook has too many cells")
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    parts.append("\t".join(values))
    finally:
        workbook.close()
    return "\n".join(parts)
