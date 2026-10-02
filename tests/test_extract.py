import io
import zipfile

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from respectscraper.extract import (
    ExtractionError,
    extract_file,
    file_type,
    normalize_url,
    parse_html,
)

LIMIT = 10 * 1024 * 1024


def minimal_pdf(text: str) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objects[3] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def test_pdf_text() -> None:
    assert "Hello PDF" in extract_file(minimal_pdf("Hello PDF"), ".pdf", max_bytes=LIMIT)


def test_docx_paragraphs_and_tables() -> None:
    doc = DocxDocument()
    doc.add_paragraph("First paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "a", "b"
    buffer = io.BytesIO()
    doc.save(buffer)

    text = extract_file(buffer.getvalue(), ".docx", max_bytes=LIMIT)

    assert "First paragraph" in text and "a\tb" in text


def test_xlsx_rows() -> None:
    book = Workbook()
    book.active.title = "Data"
    book.active.append(["name", "value"])
    book.active.append(["x", 1])
    buffer = io.BytesIO()
    book.save(buffer)

    assert extract_file(buffer.getvalue(), ".xlsx", max_bytes=LIMIT) == "# Data\nname\tvalue\nx\t1"


def test_zip_bombs_are_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * 5_000_000)

    with pytest.raises(ExtractionError, match="compression ratio"):
        extract_file(buffer.getvalue(), ".docx", max_bytes=LIMIT)


def test_text_encodings_are_detected() -> None:
    assert extract_file("café".encode(), ".txt", max_bytes=LIMIT) == "café"
    latin = "Le café est très chaud et très bon, dit Hélène à Zoé.".encode("latin-1")
    assert "très" in extract_file(latin, ".txt", max_bytes=LIMIT)


def test_failures_raise_instead_of_returning_error_text() -> None:
    with pytest.raises(ExtractionError):
        extract_file(b"not a pdf", ".pdf", max_bytes=LIMIT)
    with pytest.raises(ExtractionError, match="unsupported"):
        extract_file(b"data", ".doc", max_bytes=LIMIT)


def test_html_parsing() -> None:
    html = (
        b'<html><head><title>T</title><meta name="description" content="D">'
        b"<script>var x = 1;</script></head><body><h1>Big</h1><p>small words</p>"
        b'<a href="/a#frag">a</a><a href="/a">dup</a><a rel="nofollow" href="/n">n</a>'
        b'<a href="mailto:x@y.z">mail</a></body></html>'
    )
    page = parse_html(html, "https://Example.com/start")

    assert (page.title, page.description, page.text) == ("T", "D", "Big small words a dup n mail")
    assert page.links == ["https://example.com/a"]


def test_url_normalisation() -> None:
    assert normalize_url("HTTPS://Example.COM:443/a?b=1#c") == "https://example.com/a?b=1"
    assert normalize_url("http://example.com") == "http://example.com/"
    assert file_type("https://x.test/report.PDF?dl=1") == ".pdf"
    assert file_type("https://x.test/download", "application/pdf") == ".pdf"
